# -*- coding: utf-8 -*-
"""生成第二阶段阶段报告（Word）"""
from docx import Document
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
import os

BASE = r"C:\Users\lenovo\Doubao\chats\2026-09-11\new-chat-1\SQL-job\stage2"
SHOTS = os.path.join(BASE, "result", "screenshots")
OUT = os.path.join(BASE, "第二阶段阶段报告.docx")

doc = Document()

# 设置默认字体
style = doc.styles['Normal']
style.font.name = '宋体'
style.font.size = Pt(12)
style._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')

def add_heading(text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.name = '黑体'
        run._element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
        run.font.color.rgb = RGBColor(0, 0, 0)
    return h

def add_para(text, bold=False, size=12):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = '宋体'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    run.font.size = Pt(size)
    run.bold = bold
    p.paragraph_format.first_line_indent = Cm(0.74)
    p.paragraph_format.line_spacing = 1.5
    return p

def add_image(path, width=6.0):
    if os.path.exists(path):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run()
        run.add_picture(path, width=Inches(width))
        return p
    return None

def add_caption(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.font.name = '宋体'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    run.font.size = Pt(10)
    run.italic = True
    return p

# ========== 封面 ==========
for _ in range(4):
    doc.add_paragraph()

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title.add_run("数据库系统课程设计")
run.font.name = '黑体'
run._element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
run.font.size = Pt(28)
run.bold = True

subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = subtitle.add_run("第二阶段阶段报告")
run.font.name = '黑体'
run._element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
run.font.size = Pt(24)
run.bold = True

doc.add_paragraph()
doc.add_paragraph()

info_items = [
    ("项目名称", "小卖部管理系统"),
    ("阶段", "第二阶段（第5-9周）"),
    ("数据库", "SQL Server 2025"),
    ("Web框架", "Spring Boot 3.2"),
    ("组长", "胡海博（24325098）"),
    ("组员", "胡懿桓（24325103）、林一（24325168）"),
]
for label, value in info_items:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(f"{label}：{value}")
    run.font.name = '宋体'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    run.font.size = Pt(14)

doc.add_page_break()

# ========== 目录 ==========
add_heading("目录", level=1)
toc_items = [
    "一、设计思路",
    "  1.1 经营场景与业务流程",
    "  1.2 数据库设计（v1.0规范化）",
    "  1.3 系统分层架构",
    "二、实验过程",
    "  2.1 数据库规范化迁移",
    "  2.2 Spring Boot Web应用开发",
    "  2.3 功能演示与截图",
    "  2.4 集成测试与异常处理",
    "三、实验总结",
    "  3.1 完成情况",
    "  3.2 遇到的问题与解决",
    "  3.3 后续改进方向",
]
for item in toc_items:
    p = doc.add_paragraph(item)
    p.paragraph_format.line_spacing = 1.5

doc.add_page_break()

# ========== 一、设计思路 ==========
add_heading("一、设计思路", level=1)

add_heading("1.1 经营场景与业务流程", level=2)
add_para("本项目以校园小卖部为真实经营场景，围绕商品、库存、订单、会员四大核心业务实体构建关系数据库系统。小卖部日常经营流程包括：商品入库与库存管理、顾客选购商品并下单、订单支付与库存扣减、会员积分累计与等级管理。")
add_para("进入数据库的信息包括：商品基本信息（编号、名称、售价、进价、单位、类别、状态）、库存信息（商品编号、库存数量、货架位置、预警阈值）、订单信息（订单号、下单时间、金额、支付方式、状态、会员、收银员）、订单明细（订单号、商品编号、数量、单价、小计）、会员信息（会员编号、姓名、手机号、等级、积分）、员工信息（员工编号、姓名、账号、密码、角色）。")
add_para("暂不进入数据库的信息包括：供应商详细信息、采购订单、财务报表、员工考勤等。原因是这些信息超出了最小可运行闭环的范围，且与核心的商品-库存-订单业务流程关联较弱，留待后续迭代扩展。")

add_heading("1.2 数据库设计（v1.0规范化）", level=2)
add_para("第二阶段在v0.1基础上进行了规范化迁移。通过函数依赖分析，将原商品表中的类别名称字段拆分为独立的商品类别表，消除传递依赖（商品编号→类别编号→类别名称），达到第三范式（3NF）。")
add_para("核心表结构包括：商品表（products）、商品类别表（categories）、库存表（inventory）、订单表（orders）、订单明细表（order_items）、会员表（members）、员工表（employees）。各表主码采用业务编号（如P0001、ORD202609150001），外码建立表间关联，确保参照完整性。")

add_heading("1.3 系统分层架构", level=2)
add_para("Web应用采用经典的四层架构：页面层（Controller + Thymeleaf模板）、业务层（Service）、数据访问层（Repository/JdbcTemplate）、数据库层（SQL Server）。页面层负责用户交互与表单校验，业务层负责业务规则判断与事务管理，数据访问层负责参数化SQL执行与连接管理，数据库层负责数据持久化与约束 enforcement。")

doc.add_page_break()

# ========== 二、实验过程 ==========
add_heading("二、实验过程", level=1)

add_heading("2.1 数据库规范化迁移", level=2)
add_para("编写migration.sql脚本，完成从v0.1到v1.0的迁移：创建商品类别表、从商品表提取类别数据、调整商品表结构（类别名称改为类别编号外码）、建立外码约束。迁移前保存v0.1数据副本及关键查询结果作为基线，迁移后通过regression.sql进行回归验证，确保订单金额、商品销量、库存数量等核心指标不变。")

add_heading("2.2 Spring Boot Web应用开发", level=2)
add_para("使用Spring Boot 3.2 + Thymeleaf + Spring Security + JdbcTemplate构建Web应用。实现了员工/会员双登录模式、商品CRUD、库存查询与调整、订单创建与查询、订单明细自动计算、库存自动扣减等功能。所有数据库操作使用参数化查询，防止SQL注入。")

add_heading("2.3 功能演示与截图", level=2)

# 截图1：登录页
add_para("（1）系统登录页面", bold=True)
add_image(os.path.join(SHOTS, "01_login.png"), width=6.0)
add_caption("图1 系统登录页面（支持员工登录与会员登录切换）")

# 截图2：商品管理
add_para("（2）商品管理页面（店长视角）", bold=True)
add_image(os.path.join(SHOTS, "02_products_manager.png"), width=6.0)
add_caption("图2 商品管理页面——支持商品查询、新增、修改、删除")

# 截图3：库存管理
add_para("（3）库存管理页面", bold=True)
add_image(os.path.join(SHOTS, "03_inventory.png"), width=6.0)
add_caption("图3 库存管理页面——低库存商品自动红色预警")

# 截图4：订单管理
add_para("（4）订单管理页面", bold=True)
add_image(os.path.join(SHOTS, "04_orders.png"), width=6.0)
add_caption("图4 订单管理页面——订单列表含状态、支付方式、会员与收银员信息")

# 截图5：创建订单
add_para("（5）创建订单页面", bold=True)
add_image(os.path.join(SHOTS, "05_order_create.png"), width=6.0)
add_caption("图5 创建订单页面——选择商品、自动计算小计与合计、提交后扣减库存")

# 截图6：店员首页
add_para("（6）店员角色首页", bold=True)
add_image(os.path.join(SHOTS, "06_cashier_home.png"), width=6.0)
add_caption("图6 店员角色首页——权限受限视图")

# 截图7：越权拒绝
add_para("（7）越权访问拦截", bold=True)
add_image(os.path.join(SHOTS, "07_permission_denied.png"), width=6.0)
add_caption("图7 越权访问返回统一错误JSON——无SQL语句、堆栈等内部信息泄露")

# 截图8：未登录
add_para("（8）未登录访问拦截", bold=True)
add_image(os.path.join(SHOTS, "08_not_logged_in.png"), width=6.0)
add_caption("图8 未登录访问受保护页面——返回\"请先登录\"认证错误")

add_heading("2.4 集成测试与异常处理", level=2)
add_para("完成15项集成测试，覆盖：正常登录、商品CRUD、库存调整、订单创建与取消、参数错误、未登录访问、无权限操作、数据库约束冲突等场景。后端通过@ControllerAdvice实现全局异常处理，将参数错误、认证失败、权限不足、资源不存在、约束冲突、系统异常分别映射为统一错误码和友好提示，前端仅显示简明提示，详细错误写入后端日志。")

doc.add_page_break()

# ========== 三、实验总结 ==========
add_heading("三、实验总结", level=1)

add_heading("3.1 完成情况", level=2)
add_para("第二阶段圆满完成以下目标：")
add_para("① 数据库规范化：完成v0.1到v1.0的迁移，商品类别独立成表，达到3NF，迁移前后核心数据一致。")
add_para("② Web应用开发：基于Spring Boot构建四层架构应用，实现商品、库存、订单、订单明细的完整CRUD与业务流程。")
add_para("③ 角色权限：区分店长（manager）与店员（cashier）角色，前端隐藏操作按钮，后端业务层二次校验，关键操作使用受限数据库账户。")
add_para("④ 安全稳定：全局异常处理、参数化查询、统一响应格式、后端日志记录，确保系统安全可靠。")
add_para("⑤ 集成测试：15项测试全部通过，覆盖正常流程与异常场景。")

add_heading("3.2 遇到的问题与解决", level=2)
add_para("问题1：数据库平台从MySQL迁移到SQL Server，语法差异（反引号、ENUM、AUTO_INCREMENT等）需要逐一调整。解决：系统梳理语法差异，使用SQL Server标准语法重写所有脚本。")
add_para("问题2：Web框架从Flask切换到Spring Boot，需要重新搭建项目结构。解决：按照Spring Boot最佳实践分层组织代码，使用JdbcTemplate而非ORM以保持SQL可控性。")
add_para("问题3：登录API字段名不匹配（login_type vs loginType）。解决：通过@JsonProperty注解统一JSON字段映射。")

add_heading("3.3 后续改进方向", level=2)
add_para("① 增加采购管理模块，完善供应商与采购流程。")
add_para("② 引入Redis缓存热点数据，提升查询性能。")
add_para("③ 增加数据可视化仪表盘，展示销售趋势与库存预警。")
add_para("④ 完善单元测试覆盖率，引入CI/CD自动化部署。")

# 保存
doc.save(OUT)
print(f"报告已生成: {OUT}")
print(f"文件大小: {os.path.getsize(OUT) / 1024 / 1024:.2f} MB")
